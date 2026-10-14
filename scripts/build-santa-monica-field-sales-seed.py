#!/usr/bin/env python3
"""Build the Santa Monica field-sales seed from the researched workbook and PDF."""

from __future__ import annotations

import argparse
import json
import re
import shutil
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader


HEADER_ROW = 6


def clean(value):
    return str(value or "").strip()


def link(cell):
    return clean(cell.hyperlink.target if cell.hyperlink else "")


def useful_public_link(cell):
    value = link(cell) or clean(cell.value)
    return value if value.startswith(("http://", "https://")) else ""


def rows_by_id(sheet):
    result = {}
    for row in range(HEADER_ROW + 1, sheet.max_row + 1):
        record_id = clean(sheet.cell(row, 1).value)
        if record_id:
            result[record_id] = row
    return result


def build_seed(workbook_path: Path, pdf_path: Path, output_path: Path, photos_dir: Path):
    workbook = load_workbook(workbook_path, data_only=True)
    businesses = workbook["Businesses"]
    visits = workbook["Visit Order"]
    channels = workbook["Channels and Ads"]
    overview = workbook["Route Overview"]

    business_rows = rows_by_id(businesses)
    visit_rows = {clean(visits.cell(row, 4).value): row for row in range(HEADER_ROW + 1, visits.max_row + 1) if clean(visits.cell(row, 4).value)}
    channel_rows = rows_by_id(channels)
    route_areas = {
        clean(overview.cell(row, 1).value): clean(overview.cell(row, 2).value)
        for row in range(HEADER_ROW + 1, overview.max_row + 1)
        if clean(overview.cell(row, 1).value)
    }

    reader = PdfReader(str(pdf_path))
    photo_manifest = {}
    photos_dir.mkdir(parents=True, exist_ok=True)
    for old_file in photos_dir.glob("SM*-*"):
        if old_file.is_file():
            old_file.unlink()
    for page_number, page in enumerate(reader.pages[1:], start=2):
        text = page.extract_text() or ""
        match = re.search(r"\b(SM\d{3})\b", text)
        if not match:
            raise RuntimeError(f"No record ID found on PDF page {page_number}")
        record_id = match.group(1)
        saved = []
        for image_index, image in enumerate(page.images, start=1):
            suffix = Path(image.name).suffix.lower() or ".jpg"
            if suffix == ".jpeg":
                suffix = ".jpg"
            file_name = f"{record_id}-{image_index:02d}{suffix}"
            destination = photos_dir / file_name
            destination.write_bytes(image.data)
            saved.append({
                "url": f"/assets/field-sales-prospects/{file_name}",
                "fileName": file_name,
                "contentType": "image/png" if suffix == ".png" else "image/jpeg",
                "sourceLabel": "Public website or directory image; location unverified",
            })
        photo_manifest[record_id] = saved

    records = []
    for record_id, row in business_rows.items():
        visit_row = visit_rows.get(record_id)
        channel_row = channel_rows.get(record_id)
        route = clean(visits.cell(visit_row, 1).value) if visit_row else ""
        phone = clean(businesses.cell(row, 10).value)
        if phone.lower().startswith(("unverified", "not available", "unknown")):
            phone = ""
        channel_links = {}
        if channel_row:
            for name, column in (("instagram", 4), ("facebook", 5), ("tiktok", 6), ("xiaohongshu", 7), ("yelp", 8)):
                value = useful_public_link(channels.cell(channel_row, column))
                if value:
                    channel_links[name] = value
        records.append({
            "publicRecordId": record_id,
            "businessName": clean(businesses.cell(row, 2).value),
            "city": clean(businesses.cell(row, 3).value),
            "address": clean(businesses.cell(row, 4).value),
            "distanceKm": businesses.cell(row, 5).value,
            "distanceBand": clean(businesses.cell(row, 6).value),
            "services": clean(businesses.cell(row, 7).value),
            "filmBrands": clean(businesses.cell(row, 8).value),
            "serviceEvidence": clean(businesses.cell(row, 9).value),
            "phone": phone,
            "website": useful_public_link(businesses.cell(row, 11)),
            "googleMapsUrl": useful_public_link(businesses.cell(row, 12)),
            "directorySource": useful_public_link(businesses.cell(row, 13)),
            "serviceSources": [value.strip() for value in clean(businesses.cell(row, 14).value).splitlines() if value.strip().startswith(("http://", "https://"))],
            "lat": businesses.cell(row, 15).value,
            "lng": businesses.cell(row, 16).value,
            "researchNotes": clean(businesses.cell(row, 17).value),
            "route": route,
            "visitRegion": route_areas.get(route, "").splitlines()[-1] if route_areas.get(route) else "",
            "visitBatch": clean(visits.cell(visit_row, 2).value) if visit_row else "",
            "visitStop": visits.cell(visit_row, 3).value if visit_row else None,
            "navigationUrl": useful_public_link(visits.cell(visit_row, 14)) if visit_row else "",
            "socialLinks": channel_links,
            "referencePhotos": photo_manifest.get(record_id, []),
        })

    records.sort(key=lambda item: (item["route"], item["visitBatch"], int(item["visitStop"] or 0)))
    if len(records) != 398:
        raise RuntimeError(f"Expected 398 records, found {len(records)}")
    photo_count = sum(len(item["referencePhotos"]) for item in records)
    if photo_count != 366:
        raise RuntimeError(f"Expected 366 photos, found {photo_count}")
    if len({item["publicRecordId"] for item in records}) != len(records):
        raise RuntimeError("Duplicate record IDs found")
    missing_required = [item["publicRecordId"] for item in records if not item["businessName"] or not item["address"]]
    if missing_required:
        raise RuntimeError(f"Missing name/address for: {missing_required}")

    payload = {
        "version": "santa-monica-field-sales-2026-10-13-v1",
        "sourceWorkbook": workbook_path.name,
        "sourcePdf": pdf_path.name,
        "sourceStore": "3212 Santa Monica Blvd, Santa Monica, CA 90404",
        "recordCount": len(records),
        "photoCount": photo_count,
        "records": records,
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"records": len(records), "photos": photo_count, "output": str(output_path), "photosDir": str(photos_dir)}, ensure_ascii=False))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", required=True, type=Path)
    parser.add_argument("--pdf", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--photos-dir", required=True, type=Path)
    args = parser.parse_args()
    build_seed(args.workbook, args.pdf, args.output, args.photos_dir)


if __name__ == "__main__":
    main()
