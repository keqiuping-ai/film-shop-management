#!/usr/bin/env python3
"""Build a regional field-sales seed and extract embedded reference photos."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path
from urllib.parse import quote_plus

from openpyxl import load_workbook


HEADER_ROW = 6


def clean(value):
    return str(value or "").strip()


def useful_public_link(cell):
    target = clean(cell.hyperlink.target if cell.hyperlink else "")
    value = target or clean(cell.value)
    return value if value.startswith(("http://", "https://")) else ""


def rows_by_id(sheet, column=1):
    return {
        clean(sheet.cell(row, column).value): row
        for row in range(HEADER_ROW + 1, sheet.max_row + 1)
        if clean(sheet.cell(row, column).value)
    }


def source_center(overview):
    summary = clean(overview.cell(3, 1).value)
    match = re.search(r"Origin\s*/\s*出发点:\s*(.*?)\s*\|", summary)
    if not match:
        raise RuntimeError("Could not read the route origin from Route Overview!A3")
    return match.group(1).strip()


def normalized_identity(value):
    return re.sub(r"[^a-z0-9]+", "", clean(value).lower())


def joined_distinct(*values):
    result = []
    for value in values:
        for part in clean(value).split("\n"):
            part = part.strip()
            if part and part not in result:
                result.append(part)
    return "\n".join(result)


def merge_architectural_records(args, records, origin):
    if not args.architectural_workbook:
        return

    workbook = load_workbook(args.architectural_workbook, data_only=True)
    businesses = workbook["All Businesses"]
    visits = workbook["Visit Order"]
    channels = workbook["Channels and Ads"]
    photos = workbook["Photo Profiles"]
    business_rows = rows_by_id(businesses, 2)
    visit_rows = rows_by_id(visits, 5)
    channel_rows = rows_by_id(channels, 2)
    photo_rows = rows_by_id(photos, 2)

    for old_file in args.photos_dir.glob(f"{args.architectural_prefix}*-*"):
        if old_file.is_file():
            old_file.unlink()

    image_blobs = {}
    for image in photos._images:
        row = image.anchor._from.row + 1
        column = image.anchor._from.col + 1
        record_id = clean(photos.cell(row, 2).value)
        if not record_id.startswith(args.architectural_prefix) or column not in (6, 7):
            continue
        image_number = 1 if column == 6 else 2
        suffix = ".jpg" if str(image.format or "jpeg").lower() in ("jpg", "jpeg") else f".{str(image.format).lower()}"
        file_name = f"{record_id}-{image_number:02d}{suffix}"
        source_column = 8 if image_number == 1 else 9
        image_blobs.setdefault(record_id, []).append({
            "fileName": file_name,
            "contentType": "image/png" if suffix == ".png" else "image/jpeg",
            "sourceUrl": useful_public_link(photos.cell(row, source_column)),
            "data": image._data(),
        })

    existing_by_name = {normalized_identity(item["businessName"]): item for item in records if normalized_identity(item["businessName"])}
    existing_by_address = {normalized_identity(item["address"]): item for item in records if normalized_identity(item["address"])}
    existing_photo_hashes = {}
    for item in records:
        hashes = set()
        for photo in item.get("referencePhotos", []):
            local_path = args.photos_dir.parent.parent / clean(photo.get("url")).lstrip("/")
            if local_path.exists():
                hashes.add(hashlib.sha256(local_path.read_bytes()).hexdigest())
        existing_photo_hashes[id(item)] = hashes

    for record_id, row in business_rows.items():
        if not record_id.startswith(args.architectural_prefix):
            continue
        if not clean(businesses.cell(row, 1).value).startswith(args.architectural_region):
            continue
        business_name = clean(businesses.cell(row, 3).value)
        address = clean(businesses.cell(row, 5).value)
        visit_row = visit_rows.get(record_id)
        channel_row = channel_rows.get(record_id)
        photo_row = photo_rows.get(record_id)
        route = clean(visits.cell(visit_row, 2).value) if visit_row else ""
        target = existing_by_name.get(normalized_identity(business_name)) or existing_by_address.get(normalized_identity(address))
        architectural_services = clean(businesses.cell(row, 8).value)
        brands = clean(businesses.cell(row, 9).value)
        phone = clean(businesses.cell(row, 10).value)
        if phone.lower().startswith(("unverified", "not available", "unknown")):
            phone = ""
        website = useful_public_link(businesses.cell(row, 11))
        directory_source = useful_public_link(businesses.cell(row, 13))
        channel_links = {}
        if channel_row:
            for name, column in (("instagram", 5), ("facebook", 6), ("tiktok", 7), ("xiaohongshu", 8), ("yelp", 9)):
                value = useful_public_link(channels.cell(channel_row, column))
                if value:
                    channel_links[name] = value
        navigation_url = useful_public_link(visits.cell(visit_row, 11)) if visit_row else ""
        if not navigation_url:
            navigation_url = f"https://www.google.com/maps/dir/?api=1&origin={quote_plus(origin)}&destination={quote_plus(address)}&travelmode=driving"

        if target is None:
            target = {
                "publicRecordId": record_id,
                "businessName": business_name,
                "marketId": args.market_id,
                "city": clean(businesses.cell(row, 4).value),
                "address": address,
                "distanceKm": businesses.cell(row, 6).value,
                "distanceBand": clean(businesses.cell(row, 7).value),
                "services": architectural_services,
                "filmBrands": brands,
                "serviceEvidence": "建筑膜公开资料 / Architectural-film public research",
                "phone": phone,
                "website": website,
                "googleMapsUrl": useful_public_link(businesses.cell(row, 12)),
                "directorySource": directory_source,
                "serviceSources": [value.strip() for value in clean(businesses.cell(row, 14).value).splitlines() if value.strip().startswith(("http://", "https://"))],
                "lat": businesses.cell(row, 15).value,
                "lng": businesses.cell(row, 16).value,
                "researchNotes": clean(businesses.cell(row, 17).value),
                "route": route,
                "visitRegion": f"{args.market_label} · 建筑膜路线 {route}",
                "visitBatch": clean(visits.cell(visit_row, 3).value) if visit_row else "",
                "visitStop": visits.cell(visit_row, 4).value if visit_row else None,
                "navigationUrl": navigation_url,
                "socialLinks": channel_links,
                "photoVerification": clean(photos.cell(photo_row, 10).value) if photo_row else "",
                "referencePhotos": [],
            }
            records.append(target)
            existing_by_name[normalized_identity(business_name)] = target
            existing_by_address[normalized_identity(address)] = target
            existing_photo_hashes[id(target)] = set()
        else:
            target["services"] = joined_distinct(target.get("services"), f"建筑膜业务 / Architectural film: {architectural_services}")
            target["filmBrands"] = joined_distinct(target.get("filmBrands"), brands)
            target["researchNotes"] = joined_distinct(target.get("researchNotes"), clean(businesses.cell(row, 17).value))
            target["serviceSources"] = list(dict.fromkeys((target.get("serviceSources") or []) + [value.strip() for value in clean(businesses.cell(row, 14).value).splitlines() if value.strip().startswith(("http://", "https://"))]))
            target["phone"] = target.get("phone") or phone
            target["website"] = target.get("website") or website
            target["directorySource"] = target.get("directorySource") or directory_source
            target["socialLinks"] = {**channel_links, **(target.get("socialLinks") or {})}

        target["architecturalSourceIds"] = list(dict.fromkeys((target.get("architecturalSourceIds") or []) + [record_id]))
        target["architecturalSourceId"] = target["architecturalSourceIds"][0]
        target["architecturalServices"] = architectural_services
        target["interestedProducts"] = list(dict.fromkeys((target.get("interestedProducts") or []) + ["建筑玻璃膜"]))
        target["currentProducts"] = joined_distinct(target.get("currentProducts"), architectural_services)
        hashes = existing_photo_hashes[id(target)]
        for photo in sorted(image_blobs.get(record_id, []), key=lambda item: item["fileName"]):
            digest = hashlib.sha256(photo["data"]).hexdigest()
            if digest in hashes:
                continue
            destination = args.photos_dir / photo["fileName"]
            destination.write_bytes(photo["data"])
            target.setdefault("referencePhotos", []).append({
                "url": f"/assets/field-sales-prospects/{photo['fileName']}",
                "fileName": photo["fileName"],
                "contentType": photo["contentType"],
                "sourceLabel": "Public architectural-film website image; location unverified",
                "sourceUrl": photo["sourceUrl"],
            })
            hashes.add(digest)


def build_seed(args):
    workbook = load_workbook(args.workbook, data_only=True)
    businesses = workbook["Businesses"]
    visits = workbook["Visit Order"]
    channels = workbook["Channels and Ads"]
    photos = workbook["Photo Profiles"]
    overview = workbook["Route Overview"]

    business_rows = rows_by_id(businesses)
    visit_rows = rows_by_id(visits, 4)
    channel_rows = rows_by_id(channels)
    photo_rows = rows_by_id(photos, 4)
    route_areas = {
        clean(overview.cell(row, 1).value): clean(overview.cell(row, 2).value)
        for row in range(HEADER_ROW + 1, overview.max_row + 1)
        if clean(overview.cell(row, 1).value)
    }
    origin = source_center(overview)

    args.photos_dir.mkdir(parents=True, exist_ok=True)
    for old_file in args.photos_dir.glob(f"{args.record_prefix}*-*"):
        if old_file.is_file():
            old_file.unlink()

    photo_manifest = {record_id: [] for record_id in business_rows}
    for image in photos._images:
        row = image.anchor._from.row + 1
        column = image.anchor._from.col + 1
        record_id = clean(photos.cell(row, 4).value)
        if not record_id or column not in (9, 10):
            continue
        image_number = 1 if column == 9 else 2
        suffix = ".jpg" if str(image.format or "jpeg").lower() in ("jpg", "jpeg") else f".{str(image.format).lower()}"
        file_name = f"{record_id}-{image_number:02d}{suffix}"
        destination = args.photos_dir / file_name
        destination.write_bytes(image._data())
        source_column = 11 if image_number == 1 else 12
        photo_manifest.setdefault(record_id, []).append({
            "url": f"/assets/field-sales-prospects/{file_name}",
            "fileName": file_name,
            "contentType": "image/png" if suffix == ".png" else "image/jpeg",
            "sourceLabel": "Public website or directory image; location unverified",
            "sourceUrl": useful_public_link(photos.cell(row, source_column)),
        })

    records = []
    for record_id, row in business_rows.items():
        if not record_id.startswith(args.record_prefix):
            raise RuntimeError(f"Unexpected record ID {record_id}; expected prefix {args.record_prefix}")
        visit_row = visit_rows.get(record_id)
        channel_row = channel_rows.get(record_id)
        photo_row = photo_rows.get(record_id)
        route = clean(visits.cell(visit_row, 1).value) if visit_row else ""
        phone = clean(businesses.cell(row, 10).value)
        if phone.lower().startswith(("unverified", "not available", "unknown")):
            phone = ""
        channel_links = {}
        if channel_row:
            for name, column in (("instagram", 4), ("facebook", 5), ("tiktok", 6), ("xiaohongshu", 7), ("yelp", 8)):
                value = useful_public_link(channels.cell(channel_row, column))
                if value:
                    channel_links[name] = value
        website = useful_public_link(businesses.cell(row, 11))
        directory_source = useful_public_link(businesses.cell(row, 13))
        if visit_row:
            website = useful_public_link(visits.cell(visit_row, 19)) or website
            directory_source = useful_public_link(visits.cell(visit_row, 20)) or directory_source
        address = clean(businesses.cell(row, 4).value)
        navigation_url = useful_public_link(visits.cell(visit_row, 14)) if visit_row else ""
        if not navigation_url:
            navigation_url = f"https://www.google.com/maps/dir/?api=1&origin={quote_plus(origin)}&destination={quote_plus(address)}&travelmode=driving"
        records.append({
            "publicRecordId": record_id,
            "businessName": clean(businesses.cell(row, 2).value),
            "marketId": args.market_id,
            "city": clean(businesses.cell(row, 3).value),
            "address": address,
            "distanceKm": businesses.cell(row, 5).value,
            "distanceBand": clean(businesses.cell(row, 6).value),
            "services": clean(businesses.cell(row, 7).value),
            "filmBrands": clean(businesses.cell(row, 8).value),
            "serviceEvidence": clean(businesses.cell(row, 9).value),
            "phone": phone,
            "website": website,
            "googleMapsUrl": useful_public_link(businesses.cell(row, 12)),
            "directorySource": directory_source,
            "serviceSources": [value.strip() for value in clean(businesses.cell(row, 14).value).splitlines() if value.strip().startswith(("http://", "https://"))],
            "lat": businesses.cell(row, 15).value,
            "lng": businesses.cell(row, 16).value,
            "researchNotes": clean(businesses.cell(row, 17).value),
            "route": route,
            "visitRegion": route_areas.get(route, "").splitlines()[-1] if route_areas.get(route) else "",
            "visitBatch": clean(visits.cell(visit_row, 2).value) if visit_row else "",
            "visitStop": visits.cell(visit_row, 3).value if visit_row else None,
            "navigationUrl": navigation_url,
            "socialLinks": channel_links,
            "photoVerification": clean(photos.cell(photo_row, 13).value) if photo_row else "",
            "referencePhotos": sorted(photo_manifest.get(record_id, []), key=lambda item: item["fileName"]),
        })

    merge_architectural_records(args, records, origin)
    records.sort(key=lambda item: (item["route"], item["visitBatch"], int(item["visitStop"] or 0)))
    photo_count = sum(len(item["referencePhotos"]) for item in records)
    if len(records) != args.expected_records:
        raise RuntimeError(f"Expected {args.expected_records} records, found {len(records)}")
    if args.expected_photos >= 0 and photo_count != args.expected_photos:
        raise RuntimeError(f"Expected {args.expected_photos} photos, found {photo_count}")
    if len({item["publicRecordId"] for item in records}) != len(records):
        raise RuntimeError("Duplicate record IDs found")
    missing_required = [item["publicRecordId"] for item in records if not item["businessName"] or not item["address"] or not item["visitRegion"] or not item["visitBatch"]]
    if missing_required:
        raise RuntimeError(f"Missing name/address/region/batch for: {missing_required}")

    payload = {
        "version": args.version,
        "marketId": args.market_id,
        "marketLabel": args.market_label,
        "sourceWorkbook": args.workbook.name,
        "supplementalSourceWorkbook": args.architectural_workbook.name if args.architectural_workbook else "",
        "sourceCenter": origin,
        "recordCount": len(records),
        "photoCount": photo_count,
        "routeCount": len({item["route"] for item in records}),
        "batchCount": len({item["visitBatch"] for item in records}),
        "records": records,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"market": args.market_id, "records": len(records), "photos": photo_count, "routes": payload["routeCount"], "batches": payload["batchCount"], "output": str(args.output)}, ensure_ascii=False))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--photos-dir", required=True, type=Path)
    parser.add_argument("--market-id", required=True)
    parser.add_argument("--market-label", required=True)
    parser.add_argument("--record-prefix", required=True)
    parser.add_argument("--version", required=True)
    parser.add_argument("--expected-records", required=True, type=int)
    parser.add_argument("--expected-photos", required=True, type=int)
    parser.add_argument("--architectural-workbook", type=Path)
    parser.add_argument("--architectural-region", default="")
    parser.add_argument("--architectural-prefix", default="")
    build_seed(parser.parse_args())


if __name__ == "__main__":
    main()
