#!/usr/bin/env python3
"""Build the Las Vegas field-sales seed and extract embedded reference photos."""

from __future__ import annotations

import argparse
import json
import shutil
from pathlib import Path
from urllib.parse import quote_plus

from openpyxl import load_workbook


HEADER_ROW = 6


def clean(value):
    return str(value or "").strip()


def useful_public_link(cell):
    target = clean(cell.hyperlink.target if cell.hyperlink else "")
    value = target or clean(cell.value)
    return value if value.startswith(("http://", "https://")) else ""


def rows_by_id(sheet, column=1):
    return {
        clean(sheet.cell(row, column).value): row
        for row in range(HEADER_ROW + 1, sheet.max_row + 1)
        if clean(sheet.cell(row, column).value)
    }


def build_seed(workbook_path: Path, output_path: Path, photos_dir: Path):
    workbook = load_workbook(workbook_path, data_only=True)
    businesses = workbook["Businesses"]
    visits = workbook["Visit Order"]
    channels = workbook["Channels and Ads"]
    photos = workbook["Photo Profiles"]
    overview = workbook["Route Overview"]

    business_rows = rows_by_id(businesses)
    visit_rows = rows_by_id(visits, 4)
    channel_rows = rows_by_id(channels)
    photo_rows = rows_by_id(photos, 4)
    route_areas = {
        clean(overview.cell(row, 1).value): clean(overview.cell(row, 2).value)
        for row in range(HEADER_ROW + 1, overview.max_row + 1)
        if clean(overview.cell(row, 1).value)
    }

    photos_dir.mkdir(parents=True, exist_ok=True)
    for old_file in photos_dir.glob("LV*-*"):
        if old_file.is_file():
            old_file.unlink()

    photo_manifest = {record_id: [] for record_id in business_rows}
    for image in photos._images:
        row = image.anchor._from.row + 1
        column = image.anchor._from.col + 1
        record_id = clean(photos.cell(row, 4).value)
        if not record_id or column not in (9, 10):
            continue
        image_number = 1 if column == 9 else 2
        suffix = ".jpg" if str(image.format or "jpeg").lower() in ("jpg", "jpeg") else f".{str(image.format).lower()}"
        file_name = f"{record_id}-{image_number:02d}{suffix}"
        destination = photos_dir / file_name
        destination.write_bytes(image._data())
        source_column = 11 if image_number == 1 else 12
        source_page = useful_public_link(photos.cell(row, source_column))
        photo_manifest.setdefault(record_id, []).append({
            "url": f"/assets/field-sales-prospects/{file_name}",
            "fileName": file_name,
            "contentType": "image/png" if suffix == ".png" else "image/jpeg",
            "sourceLabel": "Public website or directory image; location unverified",
            "sourceUrl": source_page,
        })

    records = []
    for record_id, row in business_rows.items():
        visit_row = visit_rows.get(record_id)
        channel_row = channel_rows.get(record_id)
        photo_row = photo_rows.get(record_id)
        route = clean(visits.cell(visit_row, 1).value) if visit_row else ""
        phone = clean(businesses.cell(row, 10).value)
        if phone.lower().startswith(("unverified", "not available", "unknown")):
            phone = ""
        channel_links = {}
        if channel_row:
            for name, column in (("instagram", 4), ("facebook", 5), ("tiktok", 6), ("xiaohongshu", 7), ("yelp", 8)):
                value = useful_public_link(channels.cell(channel_row, column))
                if value:
                    channel_links[name] = value
        website = useful_public_link(businesses.cell(row, 11))
        directory_source = useful_public_link(businesses.cell(row, 13))
        if visit_row:
            website = useful_public_link(visits.cell(visit_row, 19)) or website
            directory_source = useful_public_link(visits.cell(visit_row, 20)) or directory_source
        address = clean(businesses.cell(row, 4).value)
        business_name = clean(businesses.cell(row, 2).value)
        navigation_url = useful_public_link(visits.cell(visit_row, 14)) if visit_row else ""
        if not navigation_url:
            navigation_url = f"https://www.google.com/maps/dir/?api=1&origin={quote_plus('3359 W Oquendo Rd, Las Vegas, NV 89118')}&destination={quote_plus(address)}&travelmode=driving"
        photo_type = clean(photos.cell(photo_row, 13).value) if photo_row else ""
        records.append({
            "publicRecordId": record_id,
            "businessName": business_name,
            "city": clean(businesses.cell(row, 3).value),
            "address": address,
            "distanceKm": businesses.cell(row, 5).value,
            "distanceBand": clean(businesses.cell(row, 6).value),
            "services": clean(businesses.cell(row, 7).value),
            "filmBrands": clean(businesses.cell(row, 8).value),
            "serviceEvidence": clean(businesses.cell(row, 9).value),
            "phone": phone,
            "website": website,
            "googleMapsUrl": useful_public_link(businesses.cell(row, 12)),
            "directorySource": directory_source,
            "serviceSources": [value.strip() for value in clean(businesses.cell(row, 14).value).splitlines() if value.strip().startswith(("http://", "https://"))],
            "lat": businesses.cell(row, 15).value,
            "lng": businesses.cell(row, 16).value,
            "researchNotes": clean(businesses.cell(row, 17).value),
            "route": route,
            "visitRegion": route_areas.get(route, "").splitlines()[-1] if route_areas.get(route) else "",
            "visitBatch": clean(visits.cell(visit_row, 2).value) if visit_row else "",
            "visitStop": visits.cell(visit_row, 3).value if visit_row else None,
            "navigationUrl": navigation_url,
            "socialLinks": channel_links,
            "photoVerification": photo_type,
            "referencePhotos": sorted(photo_manifest.get(record_id, []), key=lambda item: item["fileName"]),
        })

    records.sort(key=lambda item: (item["route"], item["visitBatch"], int(item["visitStop"] or 0)))
    if len(records) != 148:
        raise RuntimeError(f"Expected 148 records, found {len(records)}")
    photo_count = sum(len(item["referencePhotos"]) for item in records)
    if photo_count != 119:
        raise RuntimeError(f"Expected 119 photos, found {photo_count}")
    if len({item["publicRecordId"] for item in records}) != len(records):
        raise RuntimeError("Duplicate record IDs found")
    missing_required = [item["publicRecordId"] for item in records if not item["businessName"] or not item["address"] or not item["visitRegion"]]
    if missing_required:
        raise RuntimeError(f"Missing name/address/region for: {missing_required}")

    payload = {
        "version": "las-vegas-field-sales-2026-10-13-v1",
        "sourceWorkbook": workbook_path.name,
        "sourceStore": "3359 W Oquendo Rd, Las Vegas, NV 89118",
        "recordCount": len(records),
        "photoCount": photo_count,
        "routeCount": len({item["route"] for item in records}),
        "batchCount": len({item["visitBatch"] for item in records}),
        "records": records,
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"records": len(records), "photos": photo_count, "routes": payload["routeCount"], "batches": payload["batchCount"], "output": str(output_path)}, ensure_ascii=False))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--photos-dir", required=True, type=Path)
    args = parser.parse_args()
    build_seed(args.workbook, args.output, args.photos_dir)


if __name__ == "__main__":
    main()
